"""Monatsuebersicht als Excel — die sechs Spalten der bisherigen Handliste.

Ersetzt das manuelle Abtippen der Lohn-/Gehaltslisten (Spec 2026-08-12): das
System liefert genau den Block, der sich aus den Zeitdaten ergibt —
Personalnummer, Name, Urlaubsstand Vormonat/Monat, genommene Urlaubstage,
Kranktage. Alles Weitere (Vorschuss, HVV, Pfaendung, ...) bleibt Handarbeit in
der Haupt-Excel; diese Spalten kennt das System bewusst nicht.

Aufteilung wie die vier bestehenden Dateien: Firma x Beschaeftigungsart
(``hourly`` = Lohn, ``salaried`` = Gehalt).
"""

from __future__ import annotations

import calendar
import io
import uuid
from dataclasses import dataclass
from datetime import date, timedelta

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from zeiterfassung.core.holidays import is_holiday, is_weekend
from zeiterfassung.db.scoping import tenant_scoped
from zeiterfassung.models.absence import ApprovalRequest, RequestStatus, RequestType
from zeiterfassung.models.user import User

MONATE = [
    "Januar",
    "Februar",
    "März",
    "April",
    "Mai",
    "Juni",
    "Juli",
    "August",
    "September",
    "Oktober",
    "November",
    "Dezember",
]


@dataclass(frozen=True)
class OverviewRow:
    employee_number: str
    name: str
    u_tage_vormonat: float
    u_tage_monat: float
    genommen: float
    krank: float


def _month_bounds(year: int, month: int) -> tuple[date, date]:
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


def _prev_month_name(year: int, month: int) -> str:
    return MONATE[month - 2] if month > 1 else MONATE[11]


def _overlap_working_days(
    req_from: date, req_to: date, win_from: date, win_to: date, bundesland: str
) -> int:
    """Arbeitstage eines Antrags, die in das Fenster fallen.

    Ein Antrag kann ueber die Monatsgrenze laufen — fuer die Monatsspalte
    zaehlen nur die Tage IM Monat (sonst waeren Uebergaenge doppelt drin).
    """
    start, end = max(req_from, win_from), min(req_to, win_to)
    n = 0
    cur = start
    while cur <= end:
        if not is_weekend(cur) and not is_holiday(cur, bundesland):
            n += 1
        cur += timedelta(days=1)
    return n


def _sort_key(emp: User) -> tuple[int, int, str]:
    """Nach Personalnummer sortieren — numerisch, wo moeglich (5 vor 44 vor 41001)."""
    num = (emp.employee_number or "").strip()
    if num.isdigit():
        return (0, int(num), "")
    return (1, 0, num or (emp.last_name or ""))


async def build_rows(
    db: AsyncSession,
    tenant_id: uuid.UUID,
    company_id: uuid.UUID,
    employment_type: str,
    year: int,
    month: int,
    bundesland: str = "HH",
) -> list[OverviewRow]:
    """Zeilen fuer einen Monat: Urlaubsstand vorher/nachher, genommen, krank."""
    first, last = _month_bounds(year, month)

    emp_stmt = tenant_scoped(select(User), User, tenant_id).where(
        User.company_id == company_id,
        User.employment_type == employment_type,
        # Im Monat beschaeftigt: Eintritt spaetestens am Monatsende, Austritt
        # fruehestens am Monatsanfang (Ein-/Austritte mitten im Monat zaehlen).
        (User.hire_date.is_(None)) | (User.hire_date <= last),
        (User.termination_date.is_(None)) | (User.termination_date >= first),
    )
    employees = sorted((await db.execute(emp_stmt)).scalars().all(), key=_sort_key)

    rows: list[OverviewRow] = []
    for emp in employees:
        reqs = (
            (
                await db.execute(
                    tenant_scoped(select(ApprovalRequest), ApprovalRequest, tenant_id).where(
                        ApprovalRequest.employee_id == emp.id,
                        ApprovalRequest.status == RequestStatus.APPROVED,
                        ApprovalRequest.date_from <= last,
                        ApprovalRequest.date_to >= date(year, 1, 1),
                    )
                )
            )
            .scalars()
            .all()
        )

        # Urlaubs-Basis wie in der Bilanz: im Startjahr der uebernommene
        # Reststand, sonst der Jahresanspruch.
        base_year = (
            emp.counting_start_date.year
            if emp.counting_start_date
            else (emp.hire_date.year if emp.hire_date else None)
        )
        if emp.initial_vacation_remaining_days is not None and base_year == year:
            basis = float(emp.initial_vacation_remaining_days)
        else:
            basis = float(emp.vacation_entitlement_days or 0)

        vor, im_monat, krank = 0.0, 0.0, 0.0
        for r in reqs:
            if r.request_type == RequestType.VACATION:
                vor += _overlap_working_days(
                    r.date_from, r.date_to, date(year, 1, 1), first - timedelta(days=1), bundesland
                )
                im_monat += _overlap_working_days(r.date_from, r.date_to, first, last, bundesland)
            elif "SICK" in str(r.request_type):
                krank += _overlap_working_days(r.date_from, r.date_to, first, last, bundesland)

        rows.append(
            OverviewRow(
                employee_number=(emp.employee_number or "").strip(),
                name=f"{emp.last_name}, {emp.first_name}",
                u_tage_vormonat=round(basis - vor, 2),
                u_tage_monat=round(basis - vor - im_monat, 2),
                genommen=round(im_monat, 2),
                krank=round(krank, 2),
            )
        )
    return rows


def build_workbook(months: list[tuple[int, int, list[OverviewRow]]]) -> bytes:
    """Excel bauen: ein Blatt je Monat, sechs Spalten wie in der Handliste."""
    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]

    wb = Workbook()
    wb.remove(wb.active)
    for year, month, rows in months:
        ws = wb.create_sheet(title=f"{MONATE[month - 1]} {year}"[:31])
        header = [
            "Personalnummer",
            "Name",
            f"U-Tage {_prev_month_name(year, month)}",
            f"U-Tage {MONATE[month - 1]}",
            "genommene U-Tage",
            "Krank",
        ]
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            # Personalnummer als Zahl, wo moeglich — sonst verschiebt sich beim
            # Kopieren in die Haupt-Excel die Sortierung.
            pn = int(r.employee_number) if r.employee_number.isdigit() else r.employee_number
            ws.append([pn, r.name, r.u_tage_vormonat, r.u_tage_monat, r.genommen, r.krank])
        for col, width in zip("ABCDEF", (16, 30, 16, 16, 18, 10), strict=False):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
