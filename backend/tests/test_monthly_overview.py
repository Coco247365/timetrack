"""Monatsuebersicht-Export: die sechs Spalten der bisherigen Handliste.

Prueft die Rechenlogik (Urlaubsstand vorher/nachher, genommen, krank) und die
Aufteilung nach Firma x Lohn/Gehalt — inkl. Monatsgrenze (ein Antrag ueber den
Monatswechsel zaehlt nur mit seinen Tagen IM Monat).
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime

import pytest_asyncio
from sqlalchemy.ext.asyncio import AsyncSession

from zeiterfassung.models.absence import ApprovalRequest, RequestStatus, RequestType
from zeiterfassung.models.company import Company
from zeiterfassung.models.tenant import Tenant
from zeiterfassung.models.user import User
from zeiterfassung.services import monthly_overview_service as mos


@pytest_asyncio.fixture
async def company(db: AsyncSession, tenant: Tenant) -> Company:
    now = datetime.now(tz=UTC)
    c = Company(tenant_id=tenant.id, name="Alpha", is_active=True, created_at=now, updated_at=now)
    db.add(c)
    await db.flush()
    return c


async def _emp(db, tenant, company, nr, last, typ="salaried", urlaub=30.0, init=None):
    now = datetime.now(tz=UTC)
    u = User(
        tenant_id=tenant.id,
        company_id=company.id,
        employee_number=nr,
        first_name="Test",
        hashed_password="x",
        last_name=last,
        employment_type=typ,
        hire_date=date(2020, 1, 1),
        counting_start_date=date(2026, 7, 1),
        vacation_entitlement_days=urlaub,
        initial_vacation_remaining_days=init,
        contracted_hours_per_week=40.0,
        is_active=True,
        created_at=now,
        updated_at=now,
    )
    db.add(u)
    await db.flush()
    return u


async def _req(db, tenant, emp, typ, d_from, d_to):
    now = datetime.now(tz=UTC)
    r = ApprovalRequest(
        tenant_id=tenant.id,
        employee_id=emp.id,
        request_type=typ,
        status=RequestStatus.APPROVED,
        date_from=d_from,
        date_to=d_to,
        days_count=0,
        created_at=now,
        updated_at=now,
    )
    db.add(r)
    await db.flush()
    return r


async def test_columns_and_balance(db, tenant, company) -> None:
    """Basis 26 Tage, 3 im Juli genommen -> 26 vorher, 23 nachher."""
    emp = await _emp(db, tenant, company, "22018", "Muster", init=26.0)
    await _req(db, tenant, emp, RequestType.VACATION, date(2026, 7, 6), date(2026, 7, 8))

    rows = await mos.build_rows(db, tenant.id, company.id, "salaried", 2026, 7)
    assert len(rows) == 1
    r = rows[0]
    assert (r.employee_number, r.name) == ("22018", "Muster, Test")
    assert (r.u_tage_vormonat, r.u_tage_monat, r.genommen, r.krank) == (26.0, 23.0, 3.0, 0.0)


async def test_month_boundary_counts_only_days_in_month(db, tenant, company) -> None:
    """Antrag 29.07.-04.08.: im Juli nur die Juli-Arbeitstage (Mi-Fr = 3)."""
    emp = await _emp(db, tenant, company, "5", "Beispiel", init=20.0)
    await _req(db, tenant, emp, RequestType.VACATION, date(2026, 7, 29), date(2026, 8, 4))

    juli = (await mos.build_rows(db, tenant.id, company.id, "salaried", 2026, 7))[0]
    assert juli.genommen == 3.0  # 29., 30., 31.07.
    august = (await mos.build_rows(db, tenant.id, company.id, "salaried", 2026, 8))[0]
    assert august.genommen == 2.0  # 03., 04.08. (01./02. = Wochenende)
    assert august.u_tage_vormonat == 17.0  # Juli-Endstand ist August-Anfangsstand


async def test_sick_days_and_sorting_and_type_split(db, tenant, company) -> None:
    lohn = await _emp(db, tenant, company, "44", "Dobos", typ="hourly", init=27.0)
    await _emp(db, tenant, company, "5", "Beispiel", typ="hourly", init=43.0)
    await _emp(db, tenant, company, "99", "Probe", typ="salaried", init=19.0)
    await _req(db, tenant, lohn, RequestType.SICK_LEAVE, date(2026, 7, 6), date(2026, 7, 8))

    rows = await mos.build_rows(db, tenant.id, company.id, "hourly", 2026, 7)
    # nur Lohn-Mitarbeiter, sortiert nach Personalnummer (5 vor 44)
    assert [r.employee_number for r in rows] == ["5", "44"]
    assert rows[1].krank == 3.0
    assert rows[1].u_tage_monat == 27.0  # Krank zieht keinen Urlaub ab


async def test_workbook_has_sheet_per_month_with_headers(db, tenant, company) -> None:
    from openpyxl import load_workbook

    emp = await _emp(db, tenant, company, "22018", "Muster", init=26.0)
    months = []
    for y, m in ((2026, 12), (2027, 1)):
        months.append((y, m, await mos.build_rows(db, tenant.id, company.id, "salaried", y, m)))
    wb = load_workbook(io.BytesIO(mos.build_workbook(months)))

    assert wb.sheetnames == ["Dezember 2026", "Januar 2027"]
    assert [c.value for c in wb["Januar 2027"][1]] == [
        "Personalnummer",
        "Name",
        "U-Tage Dezember",
        "U-Tage Januar",
        "genommene U-Tage",
        "Krank",
    ]
    assert wb["Januar 2027"]["A2"].value == 22018  # Personalnummer als Zahl
    assert emp is not None
